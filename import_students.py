import csv
import io
import re
import unicodedata

from models import create_student, find_student_by_name

# --- Header alias mapping ---

HEADER_ALIASES = {}

_ALIAS_DEFINITIONS = {
    'nom': [
        'nom', 'nom de famille', 'last name', 'lastname', 'surname', 'family name',
        'nom etudiant', 'nom eleve',
    ],
    'prenom': [
        'prenom', 'first name', 'firstname', 'given name',
        'prenom etudiant', 'prenom eleve',
    ],
    'email': [
        'email', 'e-mail', 'mail', 'courriel', 'adresse email', 'adresse e-mail',
        'adresse mail', 'email address',
    ],
    'promotion': [
        'promotion', 'promo', 'session', 'classe', 'groupe', 'group', 'class',
        'annee', 'year',
    ],
    'nom_candidat': [
        'nom candidat',
    ],
}


def _normalize(text):
    """Lowercase, strip accents, remove punctuation."""
    text = text.strip().lower()
    # Remove accents via NFD decomposition
    text = unicodedata.normalize('NFD', text)
    text = ''.join(c for c in text if unicodedata.category(c) != 'Mn')
    # Remove punctuation (keep letters, digits, spaces)
    text = re.sub(r'[^a-z0-9 ]', '', text)
    # Collapse whitespace
    text = re.sub(r'\s+', ' ', text).strip()
    return text


# Build the lookup dict: normalized alias -> canonical field name
for field, aliases in _ALIAS_DEFINITIONS.items():
    for alias in aliases:
        HEADER_ALIASES[_normalize(alias)] = field


def _detect_columns(headers):
    """Map raw header names to canonical field names. Returns {canonical: column_index}."""
    mapping = {}
    for idx, raw_header in enumerate(headers):
        norm = _normalize(str(raw_header))
        if norm in HEADER_ALIASES:
            canonical = HEADER_ALIASES[norm]
            if canonical not in mapping:  # first match wins
                mapping[canonical] = idx
    return mapping


# --- File readers ---

def _read_csv(file_bytes):
    """Read CSV bytes, trying multiple encodings and auto-detecting delimiter."""
    for encoding in ('utf-8-sig', 'utf-8', 'latin-1'):
        try:
            text = file_bytes.decode(encoding)
            break
        except (UnicodeDecodeError, ValueError):
            continue
    else:
        raise ValueError("Impossible de décoder le fichier CSV (encodage non reconnu).")

    # Detect delimiter — exclude \t to avoid tabs inside cells being detected as separator
    try:
        sample = text[:8192]
        dialect = csv.Sniffer().sniff(sample, delimiters=',;|')
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ','

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = [row for row in reader if any(cell.strip() for cell in row)]
    return rows


def _read_xlsx(file_bytes):
    """Read XLSX bytes via openpyxl."""
    from openpyxl import load_workbook

    wb = load_workbook(filename=io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        # Skip completely empty rows
        if any(cell is not None and str(cell).strip() for cell in row):
            rows.append([str(cell).strip() if cell is not None else '' for cell in row])
    wb.close()
    return rows


def _read_file(file_bytes, filename):
    """Read CSV or XLSX file, return list of rows (each row is a list of strings)."""
    ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
    if ext == 'csv':
        return _read_csv(file_bytes)
    elif ext == 'xlsx':
        return _read_xlsx(file_bytes)
    else:
        raise ValueError(f"Format non supporté : .{ext}")


# --- Format detection ---

def _is_suivi_memoire_format(rows):
    """Detect if the file matches the 'Suivi Mémoire' format.

    Criteria: at least 3 rows, and row index 1 (the header row) contains
    a column whose normalized name matches 'nom candidat'.
    """
    if len(rows) < 3:
        return False
    headers = rows[1]
    for raw_header in headers:
        if _normalize(str(raw_header)) == 'nom candidat':
            return True
    return False


# --- Suivi Mémoire import ---

def _import_suivi_memoire(rows):
    """Import students from 'Suivi Mémoire' format.

    Row 0 = metadata (promotion in cell B1)
    Row 1 = headers
    Row 2+ = student data
    """
    # Extract promotion from cell B1 (index 1 of row 0)
    promotion = ''
    if len(rows[0]) > 1:
        promotion = rows[0][1].strip()

    # Detect columns from row 1
    headers = rows[1]
    col_map = _detect_columns(headers)

    if 'nom_candidat' not in col_map:
        raise ValueError("Colonne 'Nom candidat' introuvable dans la ligne d'en-têtes.")

    nom_candidat_idx = col_map['nom_candidat']
    email_idx = col_map.get('email')

    data_rows = rows[2:]
    imported = 0
    skipped_duplicates = 0
    skipped_errors = 0
    errors = []

    for i, row in enumerate(data_rows, start=3):  # line 3 in the file (1-indexed)
        try:
            # Read "Nom candidat" cell and split on tab
            raw_nom = row[nom_candidat_idx] if nom_candidat_idx < len(row) else ''
            parts = raw_nom.split('\t')

            if len(parts) < 2 or not parts[0].strip() or not parts[1].strip():
                skipped_errors += 1
                errors.append(f"Ligne {i} : impossible de séparer nom/prénom (pas de tabulation trouvée).")
                continue

            nom = parts[0].strip()
            prenom = parts[1].strip()

            # Email — strip trailing tabs
            email = ''
            if email_idx is not None and email_idx < len(row):
                email = row[email_idx].strip().strip('\t')

            # Check for duplicates
            if find_student_by_name(nom, prenom):
                skipped_duplicates += 1
                continue

            create_student(nom, prenom, email, promotion)
            imported += 1

        except Exception as e:
            skipped_errors += 1
            errors.append(f"Ligne {i} : {e}")

    return {
        'imported': imported,
        'skipped_duplicates': skipped_duplicates,
        'skipped_errors': skipped_errors,
        'errors': errors,
        'total_rows': len(data_rows),
    }


# --- Generic import (separate Nom/Prénom columns) ---

def _import_generic(rows):
    """Import students from a generic CSV/XLSX with separate Nom and Prénom columns."""
    if len(rows) < 2:
        raise ValueError("Le fichier est vide ou ne contient qu'une ligne d'en-tête.")

    headers = rows[0]
    data_rows = rows[1:]
    col_map = _detect_columns(headers)

    if 'nom' not in col_map or 'prenom' not in col_map:
        found = ', '.join(f'"{h}"' for h in headers if str(h).strip())
        raise ValueError(
            f"Colonnes 'Nom' et 'Prénom' introuvables. "
            f"Colonnes détectées : {found}"
        )

    imported = 0
    skipped_duplicates = 0
    skipped_errors = 0
    errors = []

    for i, row in enumerate(data_rows, start=2):  # line numbers starting at 2 (1-indexed, after header)
        try:
            nom = row[col_map['nom']].strip() if col_map['nom'] < len(row) else ''
            prenom = row[col_map['prenom']].strip() if col_map['prenom'] < len(row) else ''
            email = row[col_map['email']].strip() if 'email' in col_map and col_map['email'] < len(row) else ''
            promotion = row[col_map['promotion']].strip() if 'promotion' in col_map and col_map['promotion'] < len(row) else ''

            if not nom or not prenom:
                skipped_errors += 1
                errors.append(f"Ligne {i} : nom ou prénom vide.")
                continue

            if find_student_by_name(nom, prenom):
                skipped_duplicates += 1
                continue

            create_student(nom, prenom, email, promotion)
            imported += 1

        except Exception as e:
            skipped_errors += 1
            errors.append(f"Ligne {i} : {e}")

    return {
        'imported': imported,
        'skipped_duplicates': skipped_duplicates,
        'skipped_errors': skipped_errors,
        'errors': errors,
        'total_rows': len(data_rows),
    }


# --- Main import function ---

def import_students_from_file(file_bytes, filename):
    """
    Parse a CSV or XLSX file and import students.

    Automatically detects the 'Suivi Mémoire' format (promotion in row 1,
    combined Nom+Prénom in a single 'Nom candidat' column separated by tab).
    Falls back to generic format (separate Nom/Prénom columns).

    Returns a dict:
        imported: int
        skipped_duplicates: int
        skipped_errors: int
        errors: list of str (detail messages)
        total_rows: int
    """
    rows = _read_file(file_bytes, filename)

    if _is_suivi_memoire_format(rows):
        return _import_suivi_memoire(rows)
    else:
        return _import_generic(rows)
