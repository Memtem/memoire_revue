import os

from docx import Document
from pypdf import PdfReader


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in {'docx', 'pdf'}


def extract_text(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    if ext == '.docx':
        return extract_docx(filepath)
    elif ext == '.pdf':
        return extract_pdf(filepath)
    else:
        raise ValueError(f"Format non supporté : {ext}")


def extract_docx(filepath):
    doc = Document(filepath)
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(' | '.join(cells))
    return '\n\n'.join(paragraphs)


def extract_pdf(filepath):
    reader = PdfReader(filepath)
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text.strip())
    return '\n\n'.join(pages)
